"""Running tests on the small-gear torque-loop rig (L1): the workbook
``70_input/72_Kleingetriebepruefstand_KV3/kst-E_Laufversuche_KV3.xlsx`` (sheet ``Tabelle1``) read
row by row into ``TestRun`` records, the remarks column split into events with their load cycle
count (re-adjustments, inspection findings, failure), the level key, the speeds and a guess of
the failure criterion derived, and an anomaly list the user confirms.

Column map (header row of the sheet, read 2026-10-18): A run id (``kst-E_<n>``), B date, C rig
(KV4 = the campaign of 2023, KV3 = 2026), D material, E pinion ME number, F wheel ME number, G oil
temperature in °C at the measuring point before the test gearing (the user, 2026-10-18: a single
temperature is always that point; ``TA`` in a remark is the aggregate, ``TP`` the point before
the gearing), H target tooth temperature, I an unnamed class column of 2023, J failure mode, K
target pinion torque in Nm, L the clamped torque, M the pinion torque after the start, N after
15 min, O after 60 min, P re-adjusted, Q wheel torque in Nm (K × 52/51), R speed in 1/min, S load
cycles, T remarks, U notes.

Rules the user fixed on 2026-10-18:

- the level key is the wheel torque of column Q rounded to whole Nm. The campaign of 2023
  (MA-0296) names its levels by the pinion torque, so its "16 Nm" is 16,3 Nm at the wheel; the
  campaign of 2026 names them by the wheel torque (15,7 Nm at the pinion). Figures use the wheel
  torque of column Q.
- speeds: on KV3 column R holds the pinion speed (3058 = 3000 1/min at the wheel, 2294 = 2250),
  on KV4 the wheel speed (2250 1/min as the reference of MA-0296, single runs at 1000, 3000 and
  4000). ``speed_wheel_rpm`` is the figure of merit, the pinion turns 52/51 times faster.
- failure criteria are kept apart: total failure (loss of drive), the stop of the rig at a torque
  loss (2026: 1 Nm; 2023: not documented, presumably 5 Nm), the substitute criterion of MA-0296
  (more than 1 Nm over 100 000 load cycles) and the failure noise. ``criterion_guess`` reads the
  columns J, T and U; the user confirms every run.
- the first macroscopic crack (the failure of the method after Stommel et al.) is bracketed from
  the remarks: ``crack_window`` gives the last inspection that states no crack and the first
  that states one. A re-adjustment without a finding is no statement.
- flags: ``estimated_cycles`` (speed sensor failure, the count is an estimate from the running
  time), ``stepped`` (kst-E_22 and the temperature matrix kst-E_63, continued at 12 Nm; excluded
  from single-level statistics), ``dry_room_temperature``, ``oil_45C`` (the first runs of 2023),
  ``lubrication_deficit_possible`` (the note of column U on the changed sensor and nozzle),
  ``planned_only`` (rows without date and cycles), ``duplicate_rows``, ``id_typo``.
"""

from __future__ import annotations

import datetime as dt
import math
import re
from collections.abc import Sequence
from pathlib import Path
from typing import Literal

import openpyxl

from gearcore.errors import InputRangeError, ParseError
from gearcore.models.common import FrozenModel

SHEET = "Tabelle1"
RATIO_WHEEL_TO_PINION = 52.0 / 51.0
"""Wheel torque over pinion torque and pinion speed over wheel speed (z_2 / z_1 of kst-E)."""

Observation = Literal[
    "readjust", "check", "no_crack", "first_signs", "cracks", "wear", "pitting", "noise", "break"
]
OBSERVATION_LABELS: dict[str, str] = {
    "readjust": "nachgespannt",
    "check": "Kontrolle ohne Nachspannen",
    "no_crack": "kein Anriss",
    "first_signs": "erste Anzeichen",
    "cracks": "Anrisse",
    "wear": "Verschleiß",
    "pitting": "Grübchen",
    "noise": "Versagensgeräusch",
    "break": "Bruch",
}
_PATTERNS: list[tuple[Observation, re.Pattern[str]]] = [
    ("no_crack", re.compile(
        r"keine (?:Anrisse|Risse|Haarrisse)|noch ohne|keine Auffälligkeit", re.IGNORECASE)),
    ("first_signs", re.compile(
        r"erste Anzeichen|leichte Haarrisse|leichte Ermüdung|Anrisse einzeln|Haarrisse|\bHR\b",
        re.IGNORECASE)),
    ("cracks", re.compile(
        r"Anrisse|\bRisse\b|Ermüdung|Fußbruchanbahnung|kommenden Zahnfußbruch|durchgängig",
        re.IGNORECASE)),
    ("wear", re.compile(r"Verschlei|heller|abrasiv", re.IGNORECASE)),
    ("pitting", re.compile(r"Grübchen", re.IGNORECASE)),
    ("noise", re.compile(r"Geräusch|laut\b|lauter", re.IGNORECASE)),
    ("break", re.compile(
        r"\bZFB\b|Zahnfußbruch|Zahnbruch|Abscheren|Scherbruch|abgebrochen|umgeknickt|zahnlos"
        r"|Zähne ab|Zahn ab|gebrochen|1\.Bruch|kaputte Zähne|\bBruch\b|Flankenbruch",
        re.IGNORECASE)),
    ("readjust", re.compile(r"nachgespannt|nachverspannt|nachverspant|nachgezogen|nachgez\.",
                            re.IGNORECASE)),
    ("check", re.compile(r"nicht nachverspannt|Bilder aufgenommen|visuelle Prüfung|Kontrolle",
                         re.IGNORECASE)),
]
_CYCLES = re.compile(r"(?<![\d,.])(\d{1,3}(?:[ .]\d{3})+|\d{5,})(?!,?\d)")
_MILLIONS = re.compile(r"(\d+(?:,\d+)?)\s*Mio\b")
_TORQUE = re.compile(r"(\d+(?:[,.]\d+)?)(?:\s*(?:-|–|\.\.\.?|bis)\s*(\d+(?:[,.]\d+)?))?\s*Nm")
_TA = re.compile(r"\bTA\s*(\d+(?:,\d+)?)\s*°?C")
_TP = re.compile(r"\bTP\s*(\d+(?:,\d+)?)\s*°?C")
_BEI = re.compile(r"\bbei\s*(\d+(?:,\d+)?)\s*°C")
_LOSS = re.compile(r"ca\.?\s*(\d+(?:,\d+)?)\s*(?:Nm\s*)?Verlust über\s*(\d[\d .]*)")
_FLANK_2H = re.compile(r"nach\s*2\s*h\s*(\d+(?:,\d+)?)\s*°C")
_ESTIMATED = re.compile(r"Schätzung|geschätzt", re.IGNORECASE)
_RUN_ID = re.compile(r"kst-E[_-](\d+)", re.IGNORECASE)


def _number(text: str) -> float:
    return float(text.replace(" ", "").replace(",", "."))


def _cycles_token(text: str) -> int:
    return int(text.replace(" ", "").replace(".", ""))


def _millions_token(text: str) -> int:
    """``"2"`` or ``"2,5"`` (million) as a count, built from the digits (no arithmetic)."""
    whole, _, fraction = text.partition(",")
    digits = whole + (fraction + "000000")[:6]
    return int(digits)


class RemarkEvent(FrozenModel):
    """One segment of the remarks column: a load cycle count with what was noted there."""

    cycles: int | None
    torques_Nm: tuple[float, ...]
    """Every torque token of the segment in order (before and after a re-adjustment)."""
    torque_after_Nm: tuple[float, float] | None
    """The range set after a re-adjustment (``auf 7,8-7,9 Nm``)."""
    aggregate_temperature_C: float | None
    gearing_temperature_C: float | None
    loss_Nm: float | None
    loss_over_cycles: int | None
    observations: tuple[Observation, ...]
    text: str
    """The segment verbatim."""

    def has(self, *observations: Observation) -> bool:
        return any(o in self.observations for o in observations)


class Stage(FrozenModel):
    """A continuation row of a stepped run (the temperature matrix kst-E_63)."""

    row: int
    date: str | None
    target_pinion_torque_Nm: float | None
    wheel_torque_Nm: float | None
    speed_raw: str | None
    speed_wheel_rpm: float | None
    cycles: int | None
    cycles_raw: str | None
    flank_temperature_2h_C: float | None
    remark: str | None


class TestRun(FrozenModel):
    run_id: str
    rows: tuple[int, ...]
    """Sheet rows (1-based) that belong to the run."""
    date: str | None
    rig: str
    material: str | None
    pinion_me: str | None
    wheel_me: str | None
    oil_temperature_C: float | None
    oil_temperature_raw: str | None
    target_tooth_temperature_raw: str | None
    class_raw: str | None
    failure_mode: str | None
    target_pinion_torque_Nm: float | None
    wheel_torque_Nm: float | None
    torques_raw: tuple[tuple[str, str], ...]
    """Columns L to P verbatim (header letter, cell text)."""
    speed_raw: str | None
    speed_pinion_rpm: float | None
    speed_wheel_rpm: float | None
    cycles: int | None
    cycles_raw: str | None
    remark: str | None
    note: str | None
    events: tuple[RemarkEvent, ...]
    stages: tuple[Stage, ...]
    flags: tuple[str, ...]
    criterion_confirmed: str | None = None
    """The criterion the user confirmed (``test_runs_user.yaml``)."""
    user_note: str | None = None

    @property
    def criterion(self) -> str:
        """The confirmed criterion, else the guess from the columns."""
        return self.criterion_confirmed or self.criterion_guess

    @property
    def cycles_wheel(self) -> float | None:
        """Load cycles of a wheel tooth: the counter of the rig counts the revolutions of the
        pinion shaft (user, 2026-10-18), so a wheel tooth sees 51/52 of the count. The tables
        carry the count; the conversion shifts lg N by 0,008 only."""
        return None if self.cycles is None else self.cycles / RATIO_WHEEL_TO_PINION

    @property
    def level_Nm(self) -> int | None:
        """The nominal torque level of the test plan (whole Nm); the wheel torque itself stays a
        float beside the key."""
        if self.wheel_torque_Nm is None or self.wheel_torque_Nm <= 0.0:
            return None
        return math.floor(self.wheel_torque_Nm + 0.5)

    @property
    def crack_window(self) -> tuple[int | None, int | None]:
        """(last count stated without a crack, first count stated with one)."""
        lower: int | None = None
        for event in self.events:
            if event.cycles is None:
                continue
            if event.has("first_signs", "cracks"):
                return lower, event.cycles
            if event.has("no_crack"):
                lower = event.cycles
        return lower, None

    @property
    def criterion_guess(self) -> str:
        """Read from the columns J, T and U; the user confirms it per run."""
        mode = (self.failure_mode or "").lower()
        texts = " ".join(e.text for e in self.events).lower()
        broke = any(e.has("break") for e in self.events) or any(
            k in mode for k in ("zfb", "scher", "zähne ab", "zahn ab", "bruch"))
        noise = any(e.has("noise") for e in self.events)
        if "estimated_cycles" in self.flags:
            return "Versagensgeräusch, Lastspielzahl geschätzt" if noise else "geschätzt"
        if broke:
            return "Totalausfall"
        if noise:
            return "Versagensgeräusch"
        if "drehmomentschwelle" in texts or "abbruch" in texts:
            return "Momentverlust"
        return "nicht dokumentiert"


def classify(text: str) -> tuple[Observation, ...]:
    """Observation tags of one remark segment (``no_crack`` suppresses ``cracks``)."""
    found: list[Observation] = []
    for tag, pattern in _PATTERNS:
        if pattern.search(text):
            found.append(tag)
    if "no_crack" in found:
        found = [t for t in found if t not in ("cracks", "first_signs")]
    if "check" in found and "readjust" in found and re.search(
            r"nicht nachverspannt", text, re.IGNORECASE):
        found.remove("readjust")
    return tuple(found)


def _segments(remark: str) -> list[str]:
    parts: list[str] = []
    for line in remark.replace("\r", "").split("\n"):
        # a stop written after a full stop ("Bilder aufgenommen. 2630000, 11,6-11,8 ...") is a
        # segment of its own
        for piece in re.split(r";|(?<=\.)\s+(?=\d{5,}\b)", line):
            piece = piece.strip(" ,")
            if not piece:
                continue
            # "(ca. 1,2 Nm Verlust über 220000), Bilder aufgenommen ..." belongs to the stop
            # before, unless the piece carries a count of its own ("2630000, 11,6-11,8 ... (ca.
            # 0,6 Verlust über 100000)")
            rest = _LOSS.sub("", piece)
            if parts and _LOSS.search(piece) and not _CYCLES.search(rest) \
                    and not _MILLIONS.search(rest):
                parts[-1] += "; " + piece
                continue
            # lists of counts such as "nachgespannt bei 1 600 000 / 3 375 000 / 4 700 000"
            if "/" in piece and re.fullmatch(r"[\d /.]+", piece.split("bei", 1)[-1].strip()):
                head, _, tail = piece.partition("bei")
                parts.extend(f"{head.strip()} bei {p.strip()}" for p in tail.split("/") if p.strip())
            else:
                parts.append(piece)
    return parts


def parse_remark(remark: str | None) -> tuple[RemarkEvent, ...]:
    """Every segment of the remarks column as an event; a segment that is only a number inherits
    the observations of the segment before it."""
    if not remark or not remark.strip():
        return ()
    events: list[RemarkEvent] = []
    previous: tuple[Observation, ...] = ()
    for segment in _segments(remark):
        counts = [_cycles_token(m.group(1)) for m in _CYCLES.finditer(segment)]
        counts += [_millions_token(m.group(1)) for m in _MILLIONS.finditer(segment)]
        torques: list[float] = []
        after: tuple[float, float] | None = None
        loss = _LOSS.search(segment)
        loss_span = loss.span() if loss else (-1, -1)
        for m in _TORQUE.finditer(segment):
            if loss_span[0] <= m.start() < loss_span[1]:
                continue  # "ca. 1,2 Nm Verlust" is a loss, not a torque reading
            first = _number(m.group(1))
            torques.append(first)
            if m.group(2) is not None:
                torques.append(_number(m.group(2)))
            if re.search(r"auf\s*$", segment[: m.start()]):
                second = _number(m.group(2)) if m.group(2) is not None else first
                after = (first, second)
        ta = _TA.search(segment)
        tp = _TP.search(segment)
        bei = _BEI.search(segment)
        observations = classify(segment)
        if not observations and not re.search(r"[A-Za-zÄÖÜäöüß]", segment):
            observations = previous
        else:
            previous = observations
        events.append(RemarkEvent(
            cycles=counts[0] if counts else None,
            torques_Nm=tuple(torques),
            torque_after_Nm=after,
            aggregate_temperature_C=_number(ta.group(1)) if ta else None,
            gearing_temperature_C=(
                _number(tp.group(1)) if tp else (_number(bei.group(1)) if bei and not ta else None)
            ),
            loss_Nm=_number(loss.group(1)) if loss else None,
            loss_over_cycles=_cycles_token(loss.group(2).strip()) if loss else None,
            observations=observations,
            text=segment,
        ))
    return tuple(events)


def _text(value: object) -> str | None:
    if value is None:
        return None
    if isinstance(value, (dt.datetime, dt.date)):
        return value.strftime("%Y-%m-%d")
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    text = str(value).strip()
    return text or None


def _float(value: object) -> float | None:
    if isinstance(value, (int, float)):
        return float(value)
    text = _text(value)
    if text is None:
        return None
    match = re.match(r"^\s*(\d+(?:[,.]\d+)?)", text)
    return _number(match.group(1)) if match else None


def _cycles(value: object) -> int | None:
    if isinstance(value, (int, float)):
        return int(value) if value > 0 else None
    text = _text(value)
    if text is None:
        return None
    match = re.search(r"\d[\d .]*", text)
    if match is None:
        return None
    digits = match.group(0).replace(" ", "").replace(".", "")
    return int(digits) if digits else None


def _speeds(rig: str, speed_raw: str | None) -> tuple[float | None, float | None]:
    """(pinion, wheel) in 1/min from column R and the rig's convention."""
    value = _float(speed_raw)
    if value is None:
        return None, None
    if rig.upper() == "KV3":
        return value, value / RATIO_WHEEL_TO_PINION
    return value * RATIO_WHEEL_TO_PINION, value


def _stage(row: int, cells: Sequence[object], rig: str) -> Stage:
    remark = _text(cells[19]) if len(cells) > 19 else None
    flank = _FLANK_2H.search(remark or "")
    _, wheel_rpm = _speeds(rig, _text(cells[17]))
    return Stage(
        row=row,
        date=_text(cells[1]),
        target_pinion_torque_Nm=_float(cells[10]),
        wheel_torque_Nm=_float(cells[16]),
        speed_raw=_text(cells[17]),
        speed_wheel_rpm=wheel_rpm,
        cycles=_cycles(cells[18]),
        cycles_raw=_text(cells[18]),
        flank_temperature_2h_C=_number(flank.group(1)) if flank else None,
        remark=remark,
    )


def _run(run_id: str, rows: Sequence[tuple[int, Sequence[object]]], flags: Sequence[str]) -> TestRun:
    row, cells = rows[0]
    cells = list(cells) + [None] * (21 - len(cells))
    rig = _text(cells[2]) or "?"
    remark = _text(cells[19])
    note = _text(cells[20])
    all_flags = list(flags)
    cycles_raw = _text(cells[18])
    cycles = _cycles(cells[18])
    if remark and _ESTIMATED.search(remark) and cycles is not None:
        all_flags.append("estimated_cycles")
    if cycles_raw and cycles is not None and not re.fullmatch(r"[\d .]+", cycles_raw):
        all_flags.append("cycles_with_text")
    material = _text(cells[3])
    if (_text(cells[9]) or "").lower() == "trocken" or (_text(cells[7]) or "").upper() == "RT":
        all_flags.append("dry_room_temperature")
    oil = _float(cells[6])
    if oil is not None and oil < 60.0:
        all_flags.append("oil_45C")
    if note and re.search(r"zu wenig Volumenstrom", note, re.IGNORECASE):
        all_flags.append("lubrication_deficit_possible")
    if (_text(cells[7]) or "").lower().startswith("(stufg") or len(rows) > 1:
        all_flags.append("stepped")
    if cycles is None and _text(cells[1]) is None:
        all_flags.append("planned_only")
    elif cycles is None:
        all_flags.append("no_cycles")
    speed_pinion, speed_wheel = _speeds(rig, _text(cells[17]))
    wheel_torque = _float(cells[16])
    target_pinion = _float(cells[10])
    if (wheel_torque is None or wheel_torque <= 0.0) and target_pinion:
        wheel_torque = target_pinion * RATIO_WHEEL_TO_PINION
        all_flags.append("wheel_torque_from_pinion")
    stages: tuple[Stage, ...] = ()
    if len(rows) > 1:
        # a stepped run: every row is a stage, the remarks of all rows give the events
        stages = tuple(_stage(r, list(c) + [None] * (21 - len(c)), rig) for r, c in rows)
        remark = "\n".join(s.remark for s in stages if s.remark) or None
        last = next((s.cycles for s in reversed(stages) if s.cycles is not None), None)
        if last is not None:
            cycles, cycles_raw = last, f"{last} (letzte Stufenzeile)"
            all_flags = [f for f in all_flags if f not in ("no_cycles", "planned_only")]
    return TestRun(
        run_id=run_id,
        rows=tuple(r for r, _ in rows),
        date=_text(cells[1]),
        rig=rig,
        material=material.strip() if material else None,
        pinion_me=_text(cells[4]),
        wheel_me=_text(cells[5]),
        oil_temperature_C=oil,
        oil_temperature_raw=_text(cells[6]),
        target_tooth_temperature_raw=_text(cells[7]),
        class_raw=_text(cells[8]),
        failure_mode=_text(cells[9]),
        target_pinion_torque_Nm=target_pinion,
        wheel_torque_Nm=wheel_torque,
        torques_raw=tuple(
            (letter, _text(cells[i]) or "") for letter, i in zip("LMNOP", range(11, 16), strict=True)
        ),
        speed_raw=_text(cells[17]),
        speed_pinion_rpm=speed_pinion,
        speed_wheel_rpm=speed_wheel,
        cycles=cycles,
        cycles_raw=cycles_raw,
        remark=remark,
        note=note,
        events=parse_remark(remark),
        stages=stages,
        flags=tuple(dict.fromkeys(all_flags)),
    )


class Workbook(FrozenModel):
    path: str
    runs: tuple[TestRun, ...]
    anomalies: tuple[str, ...]

    def run(self, run_id: str) -> TestRun:
        for run in self.runs:
            if run.run_id == run_id:
                return run
        raise InputRangeError(f"no run {run_id!r} in {self.path}")


def _normalise_id(raw: str) -> tuple[str, bool]:
    match = _RUN_ID.fullmatch(raw.strip())
    if match is None:
        raise ParseError(f"run id {raw!r} does not read as kst-E_<n>")
    return f"kst-E_{int(match.group(1))}", raw.strip() != f"kst-E_{int(match.group(1))}"


def read_rows(rows: Sequence[Sequence[object]], *, path: str = "<rows>") -> Workbook:
    """Rows of the sheet (the header first) → runs; continuation rows (empty id, same wheel) are
    stages of the run before; several rows with one id are merged, the dated one leads."""
    if not rows:
        raise ParseError(f"{path}: no rows")
    header = [_text(c) or "" for c in rows[0]]
    if not header or header[0] != "A" or not header[1].startswith("Datum"):
        raise ParseError(f"{path}: unexpected header {header[:3]!r}")
    grouped: dict[str, list[tuple[int, Sequence[object]]]] = {}
    typo: set[str] = set()
    order: list[str] = []
    current: str | None = None
    for number, cells in enumerate(rows[1:], start=2):
        cells = list(cells) + [None] * (21 - len(cells))
        if all(c is None or (isinstance(c, str) and not c.strip()) for c in cells):
            continue
        raw_id = _text(cells[0])
        if raw_id is None:
            if current is None:
                raise ParseError(f"{path} row {number}: continuation row before any run")
            grouped[current].append((number, cells))
            continue
        run_id, is_typo = _normalise_id(raw_id)
        if is_typo:
            typo.add(run_id)
        if run_id not in grouped:
            grouped[run_id] = []
            order.append(run_id)
        grouped[run_id].append((number, cells))
        current = run_id
    runs: list[TestRun] = []
    anomalies: list[str] = []
    for run_id in order:
        entries = grouped[run_id]
        flags: list[str] = []
        if run_id in typo:
            flags.append("id_typo")
        dated = [e for e in entries if _text(e[1][1]) is not None]
        continuation = [e for e in entries if _text(e[1][0]) is None]
        if len(entries) > 1 and not continuation:
            # the same id on several rows: planned rows without date and data are dropped
            flags.append("duplicate_rows")
            lead = dated[0] if dated else entries[0]
            others = [e for e in entries if e is not lead]
            kept = [lead]
            for number, cells in others:
                cyc = _cycles(cells[18])
                if cyc is not None and cyc != _cycles(lead[1][18]):
                    anomalies.append(
                        f"{run_id}: Zeile {number} nennt {cyc} Lastwechsel, Zeile {lead[0]} "
                        f"{_cycles(lead[1][18])}; die datierte Zeile führt")
            entries = kept
            if len(dated) > 1:
                anomalies.append(f"{run_id}: mehrere datierte Zeilen {[e[0] for e in dated]}")
        elif continuation:
            lead = entries[0]
            entries = [lead] + continuation
        run = _run(run_id, entries, flags)
        runs.append(run)
    for run in runs:
        for flag in run.flags:
            if flag in ("estimated_cycles", "cycles_with_text", "stepped", "id_typo",
                        "duplicate_rows", "lubrication_deficit_possible",
                        "wheel_torque_from_pinion"):
                anomalies.append(f"{run.run_id}: {flag} ({run.cycles_raw or '-'})")
        start = _float(dict(run.torques_raw).get("M"))
        if start is not None and run.wheel_torque_Nm and run.rig.upper() == "KV3":
            start_wheel = start * RATIO_WHEEL_TO_PINION
            if start_wheel > run.wheel_torque_Nm + 1.0:
                anomalies.append(
                    f"{run.run_id}: Moment nach dem Start {start_wheel:.1f} Nm am Rad gegen "
                    f"{run.wheel_torque_Nm:.1f} Nm der Stufe (Spalte M {start:.1f} Nm am Ritzel)"
                    .replace(".", ","))
        if run.level_Nm is None and "planned_only" not in run.flags:
            anomalies.append(f"{run.run_id}: kein Radmoment in Spalte Q (keine Laststufe)")
        if run.speed_wheel_rpm is None and run.cycles is not None:
            anomalies.append(f"{run.run_id}: keine Drehzahl in Spalte R")
        counts = [e.cycles for e in run.events if e.cycles is not None]
        if any(b < a for a, b in zip(counts, counts[1:], strict=False)):
            anomalies.append(f"{run.run_id}: Lastspielzahlen in Spalte T nicht aufsteigend {counts}")
        if run.cycles is not None and counts and max(counts) > run.cycles:
            anomalies.append(
                f"{run.run_id}: Spalte T nennt {max(counts)} Lastwechsel, Spalte S {run.cycles}")
    return Workbook(path=path, runs=tuple(runs), anomalies=tuple(anomalies))


def load_workbook(path: Path) -> Workbook:
    book = openpyxl.load_workbook(path, data_only=True, read_only=True)
    if SHEET not in book.sheetnames:
        raise ParseError(f"{path}: no sheet {SHEET!r}")
    rows = [tuple(r) for r in book[SHEET].iter_rows(values_only=True)]
    return read_rows(rows, path=str(path))


CONFIRMATIONS = Path(__file__).resolve().parents[1] / "data" / "measurement" / "test_runs_user.yaml"


def apply_confirmations(book: Workbook, path: Path = CONFIRMATIONS) -> Workbook:
    """The user's confirmations (criterion, note, withdrawn flags per run) on top of the
    workbook; an entry for a run the workbook does not have is an error."""
    import yaml

    data = yaml.safe_load(path.read_text(encoding="utf-8")) or {}
    entries: dict[str, dict[str, object]] = dict(data.get("runs") or {})
    known = {run.run_id for run in book.runs}
    unknown = sorted(set(entries) - known)
    if unknown:
        raise InputRangeError(f"{path.name}: runs not in the workbook: {', '.join(unknown)}")
    withdrawn: dict[str, set[str]] = {}
    for run_id, entry in entries.items():
        raw = entry.get("withdraw_flags")
        withdrawn[run_id] = {str(f) for f in raw} if isinstance(raw, list) else set()
    runs: list[TestRun] = []
    for run in book.runs:
        found = entries.get(run.run_id)
        if found is None:
            runs.append(run)
            continue
        criterion = found.get("criterion")
        note = found.get("note")
        update: dict[str, object] = {
            "flags": tuple(f for f in run.flags if f not in withdrawn[run.run_id]),
            "criterion_confirmed": str(criterion) if criterion is not None else None,
            "user_note": str(note) if note is not None else None,
        }
        speed = found.get("speed_pinion_rpm")
        if isinstance(speed, (int, float)):
            # a speed the workbook does not carry, given by the user (pinion shaft)
            update["speed_raw"] = f"{speed:g} (Nutzer)"
            update["speed_pinion_rpm"] = float(speed)
            update["speed_wheel_rpm"] = float(speed) / RATIO_WHEEL_TO_PINION
        runs.append(run.model_copy(update=update))
    with_speed = {rid for rid, e in entries.items() if isinstance(e.get("speed_pinion_rpm"), (int, float))}
    anomalies = tuple(
        a for a in book.anomalies
        if not any(a.startswith(f"{rid}: {flag}") for rid, flags in withdrawn.items()
                   for flag in flags)
        and not any(a.startswith(f"{rid}: keine Drehzahl") for rid in with_speed)
    )
    return Workbook(path=book.path, runs=tuple(runs), anomalies=anomalies)


def loss_intervals(run: TestRun) -> list[tuple[int, int, float, float, float | None, float | None]]:
    """Per re-adjustment of a 2026 run: (from, to, loss in Nm, loss per million cycles, TA, TP).
    The loss is the one the remark states (``ca. 1,2 Nm Verlust über 220000``); where it states
    none, the torque set after the previous re-adjustment minus the torque before this one."""
    out: list[tuple[int, int, float, float, float | None, float | None]] = []
    previous_cycles = 0
    previous_after: float | None = None
    stated_since_readjust = False
    for event in run.events:
        if event.cycles is None:
            continue
        loss: float | None = None
        span = event.cycles - previous_cycles
        if event.loss_Nm is not None and event.loss_over_cycles:
            loss, span = event.loss_Nm, event.loss_over_cycles
            stated_since_readjust = True
        elif (event.has("readjust") and previous_after is not None and event.torques_Nm
              and not stated_since_readjust):
            # derived only where no stated interval covers the span since the last re-adjustment
            loss = previous_after - event.torques_Nm[0]
        if loss is not None and span > 0:
            out.append((event.cycles - span, event.cycles, loss, loss / span * 1.0e6,
                        event.aggregate_temperature_C, event.gearing_temperature_C))
        if event.torque_after_Nm is not None:
            previous_after = sum(event.torque_after_Nm) / 2.0
            previous_cycles = event.cycles
            stated_since_readjust = False
    return out


__all__ = [
    "CONFIRMATIONS",
    "OBSERVATION_LABELS",
    "RATIO_WHEEL_TO_PINION",
    "Observation",
    "RemarkEvent",
    "Stage",
    "TestRun",
    "Workbook",
    "apply_confirmations",
    "classify",
    "load_workbook",
    "loss_intervals",
    "parse_remark",
    "read_rows",
]
