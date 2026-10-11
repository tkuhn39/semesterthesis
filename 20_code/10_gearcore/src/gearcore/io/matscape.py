"""Stress-strain curves of the material as the METG tensile tests delivered them for Matscape.

Two workbooks of the user (``00_literatur/06_datenblaetter/Stanyl TW 200F6_T80_50%rLF.xlsx``,
``T23_50%rLF.XLSX``) hold the engineering stress-strain curves of Stanyl TW200F6 (PA46-GF30) at one
temperature and 50 % relative humidity each. Grammar (read 2026-10-11): one sheet per load
direction, named by the angle to the flow direction with a degree sign (``0°``, ``15°``, ``45°``,
``90°``); in a sheet, pairs of columns per strain rate, a header row with the rate label above
(``QS``, ``DY1``, ``DY2``, at 80 °C with the test speed in m/s and an arrow to the strain rate in
1/s), a header row ``Dehnung [%] | Spannung [Mpa]`` and the numeric rows below; a missing curve
is marked by a text in place of the header (``gibt's nicht``). Strains in %, stresses in MPa,
strain strictly increasing from 0, the curves end at break.

The parser keeps every curve as read (``StressStrainCurve``); what counts as the static
characteristic value or a strength is decided in ``gearcore.fatigue.materials``. Temperature and
humidity are read from the file name (``T80`` and ``50%rLF``) unless given; they are provenance,
not data of the sheet.
"""

from __future__ import annotations

import bisect
import re
from collections.abc import Sequence
from pathlib import Path

import openpyxl

from gearcore.errors import InputRangeError, ParseError
from gearcore.models.common import FrozenModel

_DIRECTION = re.compile(r"^\s*(\d+)\s*[°º�]?\s*$")
_TEMPERATURE = re.compile(r"T(\d+)")
_HUMIDITY = re.compile(r"(\d+)\s*%\s*rLF", re.IGNORECASE)
_RATE = re.compile(r"^(QS|DY1|DY2)\b", re.IGNORECASE)
_SPEED = re.compile(r"\[([0-9.Ee+-]+)\s*m/s\]")


class StressStrainCurve(FrozenModel):
    """One engineering stress-strain curve: strain in %, stress in MPa, strain strictly
    increasing from zero, as the test delivered it (no smoothing, no extrapolation)."""

    strain_percent: tuple[float, ...]
    stress_mpa: tuple[float, ...]

    def stress_at(self, strain_percent: float) -> float:
        """Stress at a strain inside the measured range (linear between the points)."""
        strains = self.strain_percent
        if not strains[0] <= strain_percent <= strains[-1]:
            raise InputRangeError(
                f"strain {strain_percent:g} % lies outside the curve ({strains[0]:g} to {strains[-1]:g} %)"
            )
        j = bisect.bisect_right(strains, strain_percent)
        if j >= len(strains):
            return self.stress_mpa[-1]
        i = j - 1
        x0, x1 = strains[i], strains[j]
        y0, y1 = self.stress_mpa[i], self.stress_mpa[j]
        return y0 + (y1 - y0) * (strain_percent - x0) / (x1 - x0)

    def strain_at(self, stress_mpa: float) -> float:
        """Strain on the rising branch (up to the maximum stress) at a stress between 0 and the
        maximum; ``InputRangeError`` above the maximum (the curve does not reach it)."""
        if stress_mpa < 0.0:
            raise InputRangeError("a stress on the curve is not negative")
        peak_stress, peak_strain = self.maximum
        if stress_mpa > peak_stress:
            raise InputRangeError(
                f"stress {stress_mpa:g} MPa exceeds the maximum of the curve ({peak_stress:g} MPa)"
            )
        for i in range(1, len(self.stress_mpa)):
            if self.strain_percent[i] > peak_strain:
                break
            y0, y1 = self.stress_mpa[i - 1], self.stress_mpa[i]
            if y0 <= stress_mpa <= y1 and y1 > y0:
                x0, x1 = self.strain_percent[i - 1], self.strain_percent[i]
                return x0 + (x1 - x0) * (stress_mpa - y0) / (y1 - y0)
            if stress_mpa == y0:
                return self.strain_percent[i - 1]
        return peak_strain

    @property
    def maximum(self) -> tuple[float, float]:
        """(stress, strain) at the largest stress of the curve."""
        k = max(range(len(self.stress_mpa)), key=lambda i: self.stress_mpa[i])
        return self.stress_mpa[k], self.strain_percent[k]

    @property
    def end(self) -> tuple[float, float]:
        """(stress, strain) of the last point (break)."""
        return self.stress_mpa[-1], self.strain_percent[-1]

    @property
    def initial_modulus_mpa(self) -> float:
        """Secant modulus between the first two points, in MPa (strain converted from %)."""
        return (
            (self.stress_mpa[1] - self.stress_mpa[0])
            / (self.strain_percent[1] - self.strain_percent[0])
            * 100.0
        )


class MatscapeCurve(FrozenModel):
    """A curve with its provenance: direction to the flow, strain rate label, test speed."""

    direction_deg: int
    rate: str
    """``QS``, ``DY1`` or ``DY2``, as the sheet labels them."""
    speed_m_s: float | None
    """Test speed in m/s where the label prints it (80 °C workbook), else None."""
    label: str
    """The header text of the sheet, verbatim."""
    curve: StressStrainCurve


class MatscapeWorkbook(FrozenModel):
    """Every curve of one workbook (one temperature, one humidity)."""

    file: str
    temperature_c: float
    humidity_percent: float
    curves: tuple[MatscapeCurve, ...]
    missing: tuple[tuple[int, str], ...]
    """(direction, rate) pairs the sheet marks as absent."""

    def curve(self, direction_deg: int, rate: str) -> MatscapeCurve:
        for item in self.curves:
            if item.direction_deg == direction_deg and item.rate.upper() == rate.upper():
                return item
        raise InputRangeError(
            f"no curve at {direction_deg} deg / {rate} in {self.file}; "
            f"present: {[(c.direction_deg, c.rate) for c in self.curves]}"
        )

    @property
    def directions_deg(self) -> tuple[int, ...]:
        return tuple(sorted({c.direction_deg for c in self.curves}))


def _number(value: object) -> float | None:
    if isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        return float(value)
    return None


def _direction_of_sheet(title: str) -> int | None:
    match = _DIRECTION.match(title)
    return int(match.group(1)) if match else None


def _provenance_from_name(name: str) -> tuple[float | None, float | None]:
    temperature = _TEMPERATURE.search(name)
    humidity = _HUMIDITY.search(name)
    return (
        float(temperature.group(1)) if temperature else None,
        float(humidity.group(1)) if humidity else None,
    )


def load_matscape_workbook(
    path: Path, *, temperature_c: float | None = None, humidity_percent: float | None = None
) -> MatscapeWorkbook:
    """Read every curve of the workbook. Temperature and humidity come from the file name
    (``T80``, ``50%rLF``) when not given; a name without them needs both arguments."""
    name_temperature, name_humidity = _provenance_from_name(path.name)
    temperature = temperature_c if temperature_c is not None else name_temperature
    humidity = humidity_percent if humidity_percent is not None else name_humidity
    if temperature is None or humidity is None:
        raise ParseError(
            f"{path.name}: temperature and humidity are neither in the name (T80, 50%rLF) nor given"
        )
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    curves: list[MatscapeCurve] = []
    missing: list[tuple[int, str]] = []
    try:
        for sheet in workbook.worksheets:
            direction = _direction_of_sheet(str(sheet.title))
            if direction is None:
                raise ParseError(f"{path.name}: sheet {sheet.title!r} does not name a direction")
            rows = [tuple(row) for row in sheet.iter_rows(values_only=True)]
            header_index = next(
                (
                    i
                    for i, row in enumerate(rows)
                    if any(isinstance(v, str) and v.startswith("Dehnung") for v in row)
                ),
                None,
            )
            if header_index is None:
                raise ParseError(f"{path.name}: sheet {sheet.title!r} has no 'Dehnung' header")
            header = rows[header_index]
            width = max(len(row) for row in rows)
            for col in range(0, width, 2):
                label = _rate_label(rows[:header_index], col)
                if label is None:
                    continue
                rate_match = _RATE.match(label)
                if rate_match is None:
                    raise ParseError(
                        f"{path.name}: sheet {sheet.title!r}: unknown rate label {label!r}"
                    )
                rate = rate_match.group(1).upper()
                head = header[col] if col < len(header) else None
                if not (isinstance(head, str) and head.startswith("Dehnung")):
                    missing.append((direction, rate))
                    continue
                unit = header[col + 1] if col + 1 < len(header) else None
                if not (isinstance(unit, str) and unit.lower().startswith("spannung")):
                    raise ParseError(
                        f"{path.name}: sheet {sheet.title!r}, column {col + 1}: 'Spannung' header expected"
                    )
                strains: list[float] = []
                stresses: list[float] = []
                for row in rows[header_index + 1 :]:
                    strain = _number(row[col]) if col < len(row) else None
                    stress = _number(row[col + 1]) if col + 1 < len(row) else None
                    if strain is None and stress is None:
                        continue
                    if strain is None or stress is None:
                        raise ParseError(
                            f"{path.name}: sheet {sheet.title!r}, column {col}: a row with only one number"
                        )
                    strains.append(strain)
                    stresses.append(stress)
                if len(strains) < 2:
                    raise ParseError(f"{path.name}: sheet {sheet.title!r}, column {col}: no curve")
                if strains[0] != 0.0 or any(
                    b <= a for a, b in zip(strains, strains[1:], strict=False)
                ):
                    raise ParseError(
                        f"{path.name}: sheet {sheet.title!r}, {rate}: strain must rise strictly from 0"
                    )
                speed = _SPEED.search(label)
                curves.append(
                    MatscapeCurve(
                        direction_deg=direction,
                        rate=rate,
                        speed_m_s=float(speed.group(1)) if speed else None,
                        label=label,
                        curve=StressStrainCurve(
                            strain_percent=tuple(strains), stress_mpa=tuple(stresses)
                        ),
                    )
                )
    finally:
        workbook.close()
    if not curves:
        raise ParseError(f"{path.name}: no curve found")
    return MatscapeWorkbook(
        file=path.name,
        temperature_c=temperature,
        humidity_percent=humidity,
        curves=tuple(curves),
        missing=tuple(missing),
    )


def _rate_label(rows_above: Sequence[tuple[object, ...]], col: int) -> str | None:
    """The rate label above a column pair: the last non-empty text in the pair's columns above
    the header row (the 80 °C sheets spread it over two lines)."""
    for row in reversed(rows_above):
        for c in (col, col + 1):
            cell = row[c] if c < len(row) else None
            if isinstance(cell, str) and cell.strip():
                return cell.strip()
    return None
