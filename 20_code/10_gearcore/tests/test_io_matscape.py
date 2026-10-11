"""The METG / Matscape stress-strain workbooks (``gearcore.io.matscape``)."""

from pathlib import Path

import openpyxl
import pytest

from gearcore.errors import InputRangeError, ParseError
from gearcore.io.matscape import StressStrainCurve, load_matscape_workbook

REPO = Path(__file__).resolve().parents[3]
DATASHEETS = REPO / "00_literatur" / "06_datenblaetter"


def _workbook(path: Path, *, with_missing: bool = True) -> Path:
    book = openpyxl.Workbook()
    sheet = book.active
    assert sheet is not None
    sheet.title = "0°"
    sheet.append(["QS [1.67E-05 m/s] ", None, "DY1 [0.174 m/s]\n-> 3s-1", None])
    sheet.append(["Dehnung [%]", "Spannung [Mpa]", "Dehnung [%]", "Spannung [Mpa]"])
    for i in range(0, 11):
        strain = 0.5 * i
        sheet.append([strain, min(20.0 * strain, 80.0 + 2.0 * strain), strain, 30.0 * strain])
    if with_missing:
        other = book.create_sheet("45°")
        other.append(["QS", None, "DY2", None])
        other.append(["gibt's nicht", None, "Dehnung [%]", "Spannung [Mpa]"])
        other.append([None, None, 0.0, 0.0])
        other.append([None, None, 1.0, 40.0])
    book.save(path)
    return path


def test_reads_curves_rates_directions_and_provenance(tmp_path: Path) -> None:
    path = _workbook(tmp_path / "Stanyl TW 200F6_T80_50%rLF.xlsx")
    book = load_matscape_workbook(path)
    assert book.temperature_c == 80.0 and book.humidity_percent == 50.0
    assert book.directions_deg == (0, 45)
    assert book.missing == ((45, "QS"),)
    qs = book.curve(0, "QS")
    assert qs.speed_m_s == pytest.approx(1.67e-5)
    assert qs.label.startswith("QS")
    curve = qs.curve
    assert curve.strain_percent[0] == 0.0 and len(curve.strain_percent) == 11
    assert curve.maximum == (90.0, 5.0)  # the plateau 80 + 2 strain reaches 90 at 5 %
    assert curve.stress_at(1.25) == pytest.approx(25.0)
    assert curve.strain_at(25.0) == pytest.approx(1.25)
    assert curve.initial_modulus_mpa == pytest.approx(2000.0)
    dy1 = book.curve(0, "dy1").curve
    assert dy1.stress_at(2.0) == pytest.approx(60.0)
    with pytest.raises(InputRangeError):
        book.curve(90, "QS")
    with pytest.raises(InputRangeError):
        curve.strain_at(95.0)
    with pytest.raises(InputRangeError):
        curve.stress_at(7.0)


def test_provenance_must_be_known(tmp_path: Path) -> None:
    path = _workbook(tmp_path / "curves.xlsx", with_missing=False)
    with pytest.raises(ParseError, match="temperature and humidity"):
        load_matscape_workbook(path)
    book = load_matscape_workbook(path, temperature_c=23.0, humidity_percent=50.0)
    assert book.temperature_c == 23.0 and book.missing == ()


def test_sheet_without_direction_or_header_is_refused(tmp_path: Path) -> None:
    book = openpyxl.Workbook()
    sheet = book.active
    assert sheet is not None
    sheet.title = "Daten"
    sheet.append(["Dehnung [%]", "Spannung [Mpa]"])
    path = tmp_path / "T23_50%rLF.xlsx"
    book.save(path)
    with pytest.raises(ParseError, match="direction"):
        load_matscape_workbook(path)
    book = openpyxl.Workbook()
    sheet = book.active
    assert sheet is not None
    sheet.title = "0°"
    sheet.append(["QS", None])
    sheet.append(["x", "y"])
    book.save(path)
    with pytest.raises(ParseError, match="Dehnung"):
        load_matscape_workbook(path)


def test_curve_model_is_strict() -> None:
    curve = StressStrainCurve(strain_percent=(0.0, 1.0, 2.0), stress_mpa=(0.0, 50.0, 60.0))
    assert curve.end == (60.0, 2.0)
    assert curve.strain_at(60.0) == 2.0
    with pytest.raises(InputRangeError):
        curve.strain_at(-1.0)


@pytest.mark.skipif(
    not (DATASHEETS / "T23_50%rLF.XLSX").is_file(), reason="METG workbooks not present"
)
def test_metg_workbooks_of_the_user() -> None:
    """The numbers read on 2026-10-11 (0 deg, QS): 23 degC 115,51 MPa at 6,39 %; 80 degC 88,99
    MPa at 6,98 %; the 45 deg QS curve at 80 degC is marked absent."""
    room = load_matscape_workbook(DATASHEETS / "T23_50%rLF.XLSX")
    assert room.temperature_c == 23.0 and room.directions_deg == (0, 15, 45, 90)
    peak, strain = room.curve(0, "QS").curve.maximum
    assert peak == pytest.approx(115.51, abs=0.01) and strain == pytest.approx(6.39, abs=0.01)
    hot = load_matscape_workbook(DATASHEETS / "Stanyl TW 200F6_T80_50%rLF.xlsx")
    assert hot.temperature_c == 80.0 and hot.directions_deg == (0, 45, 90)
    peak, strain = hot.curve(0, "QS").curve.maximum
    assert peak == pytest.approx(88.99, abs=0.01) and strain == pytest.approx(6.98, abs=0.01)
    assert (45, "QS") in hot.missing
    assert hot.curve(0, "DY2").speed_m_s == pytest.approx(10.0)
